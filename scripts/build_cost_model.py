from __future__ import annotations

import json
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / "智慧安防_成本选型逻辑模型_v1.xlsx"
BASE_URL = "http://localhost:4173"


def fetch_json(path: str) -> dict:
    with urllib.request.urlopen(f"{BASE_URL}{path}") as response:
        return json.load(response)


def style_cell(cell, fill=None, bold=False, color="07182B", size=11, align="left"):
    thin = Side(style="thin", color="D9E3EA")
    cell.font = Font(name="Microsoft YaHei", bold=bold, color=color, size=size)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


def setup_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        style_cell(cell, fill="07182B", bold=True, color="FFFFFF", align="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value is not None:
                style_cell(cell, fill="FFFFFF")
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def add_headers(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        style_cell(cell, fill="07182B", bold=True, color="FFFFFF", align="center")


def build_workbook():
    catalog = fetch_json("/api/catalog")
    sizing_catalog = fetch_json("/api/sizing/catalog")

    feature_costs = sizing_catalog["featureCostProfiles"]
    resolutions = sizing_catalog["resolutions"]
    hardware = sizing_catalog["hardwareProfiles"]

    wb = Workbook()
    wb.remove(wb.active)

    ws_overview = wb.create_sheet("01逻辑总览")
    ws_groups = wb.create_sheet("02摄像头组输入")
    ws_features = wb.create_sheet("03功能选择输入")
    ws_calc = wb.create_sheet("04计算明细")
    ws_summary = wb.create_sheet("05报价汇总")
    ws_feature_cost = wb.create_sheet("06算法成本画像")
    ws_res = wb.create_sheet("07分辨率倍率")
    ws_hw = wb.create_sheet("08硬件档位")
    ws_notes = wb.create_sheet("09参数说明")

    # 01 overview
    ws_overview["A1"] = "智慧安防成本与硬件选型逻辑模型"
    ws_overview["A1"].font = Font(name="Microsoft YaHei", bold=True, size=18, color="07182B")
    ws_overview["A3"] = (
        "这份表不是最终报价单，而是把“客户需求 → 算法/算力 → 硬件 → 固定成本/用量成本”的计算逻辑摊开，"
        "用于内部复盘和参数校准。"
    )
    ws_overview.merge_cells("A3:H4")
    ws_overview["A3"].alignment = Alignment(wrap_text=True)
    ws_overview.append([])
    add_headers(ws_overview, ["环节", "说明"])
    for item in [
        ("1 客户输入", "摄像头组、路数、分辨率、每组选择的识别功能"),
        ("2 算法画像", "每个功能对应算法等级、1080P单路算力、研发/适配成本、AI复核策略"),
        ("3 算力计算", "路数 × 分辨率倍率 × 单路算法算力；再加安全余量"),
        ("4 硬件选型", "推荐硬件需同时满足总算力与总路数；超过单机则按多节点估算"),
        ("5 价格拆分", "硬件买断/月付 + 功能研发/适配 + 月度算法运维；token/API另计"),
    ]:
        ws_overview.append(item)
    ws_overview.append([])
    ws_overview.append(["成本项", "口径"])
    for item in [
        ("硬件成本", "由总路数和总算力决定，可买断或月付；这是最需要保守估算的部分。"),
        ("功能研发/适配成本", "按唯一功能计一次，不按摄像头路数重复收；复杂定制功能可单独调高。"),
        ("月度算法运维成本", "按唯一功能计一次，用于模型维护、规则调优、场景样本优化等。"),
        ("AI/token费用", "不放进固定成本，只估算事件量，后续作为用量服务包或按量收费。"),
    ]:
        ws_overview.append(item)

    # 02 camera group input
    add_headers(ws_groups, ["组编号", "摄像头组名称", "摄像头路数", "分辨率", "厂商", "接入方式", "位置/备注"])
    for row in [
        ["G01", "生产车间摄像头", 48, "1080P", "海康", "RTSP", "安全帽、工服、抽烟、危险区域"],
        ["G02", "仓库摄像头", 36, "1080P", "海康", "RTSP", "仓库与消防通道"],
        ["G03", "园区周界摄像头", 28, "2K", "大华", "GB28181", "周界围栏与夜间入侵"],
        ["G04", "门岗车辆摄像头", 12, "4K", "宇视", "ONVIF", "门岗出入口"],
    ]:
        ws_groups.append(row)
    for row in range(6, 101):
        ws_groups.cell(row, 1, f"G{row - 1:02d}")

    # 03 feature selection input
    add_headers(ws_features, ["组编号", "功能ID", "功能名称自动显示", "备注"])
    for row in [
        ["G01", "helmet", "", ""],
        ["G01", "vest", "", ""],
        ["G01", "smoking", "", ""],
        ["G01", "danger-zone", "", ""],
        ["G02", "cargo-block", "", ""],
        ["G02", "channel-block", "", ""],
        ["G02", "fire-lane-block", "", ""],
        ["G02", "smoke", "", ""],
        ["G03", "perimeter-intrusion", "", ""],
        ["G03", "line-crossing", "", ""],
        ["G03", "night-intrusion", "", ""],
        ["G04", "plate", "", ""],
        ["G04", "illegal-parking", "", ""],
        ["G04", "vehicle-stay", "", ""],
    ]:
        ws_features.append(row)
    for row in range(2, 201):
        ws_features.cell(row, 3, f'=IFERROR(VLOOKUP(B{row},\'06算法成本画像\'!$A:$B,2,FALSE),"")')

    # 06 feature cost profile
    add_headers(
        ws_feature_cost,
        [
            "功能ID",
            "功能名称",
            "类别",
            "算法族",
            "算法等级",
            "1080P单路算力",
            "一次性研发/适配成本",
            "月度算法运维成本",
            "AI复核策略",
            "单路每日AI事件量",
            "说明",
        ],
    )
    for item in feature_costs:
        ws_feature_cost.append(
            [
                item["feature_id"],
                item["feature_name"],
                item["category_name"],
                item["algorithm_family"],
                item["workload_level"],
                item["base_compute_units_1080p"],
                item["one_time_dev_cost"],
                item["monthly_algorithm_ops_cost"],
                item["ai_review_policy"],
                item["ai_event_rate_per_camera_day"],
                item.get("notes", ""),
            ]
        )

    # 07 resolution profile
    add_headers(ws_res, ["分辨率", "算力倍率", "说明"])
    for item in resolutions:
        ws_res.append([item["name"], item["compute_multiplier"], item.get("notes", "")])

    # 08 hardware profile
    add_headers(
        ws_hw,
        ["硬件ID", "硬件名称", "可承载算力单位", "建议最大路数", "买断成本", "月付成本", "CPU", "GPU", "内存", "说明", "是否满足当前需求"],
    )
    for item in hardware:
        ws_hw.append(
            [
                item["id"],
                item["name"],
                item["capacity_units"],
                item["recommended_max_streams"],
                item["purchase_cost"],
                item["monthly_cost"],
                item.get("cpu", ""),
                item.get("gpu", ""),
                item.get("memory", ""),
                item.get("notes", ""),
                "",
            ]
        )
    for row in range(2, 2 + len(hardware)):
        ws_hw.cell(row, 11, f'=IF(AND(C{row}>=\'05报价汇总\'!$B$9,D{row}>=\'05报价汇总\'!$B$5),1,0)')

    # 04 calculation detail
    add_headers(
        ws_calc,
        [
            "行号",
            "组编号",
            "摄像头组名称",
            "摄像头路数",
            "分辨率",
            "功能ID",
            "功能名称",
            "分辨率倍率",
            "1080P单路算力",
            "本行算力",
            "算法等级",
            "AI复核策略",
            "单路每日AI事件量",
            "本行每日AI事件量",
            "唯一功能标记",
            "一次性研发成本(首次计)",
            "月度算法运维成本(首次计)",
            "说明",
        ],
    )
    for row in range(2, 201):
        ws_calc.cell(row, 1, row - 1)
        ws_calc.cell(row, 2, f"='03功能选择输入'!A{row}")
        ws_calc.cell(row, 3, f'=IFERROR(VLOOKUP(B{row},\'02摄像头组输入\'!$A:$G,2,FALSE),"")')
        ws_calc.cell(row, 4, f"=IFERROR(VLOOKUP(B{row},'02摄像头组输入'!$A:$G,3,FALSE),0)")
        ws_calc.cell(row, 5, f'=IFERROR(VLOOKUP(B{row},\'02摄像头组输入\'!$A:$G,4,FALSE),"")')
        ws_calc.cell(row, 6, f"='03功能选择输入'!B{row}")
        ws_calc.cell(row, 7, f'=IFERROR(VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,2,FALSE),"")')
        ws_calc.cell(row, 8, f"=IFERROR(VLOOKUP(E{row},'07分辨率倍率'!$A:$C,2,FALSE),1)")
        ws_calc.cell(row, 9, f"=IFERROR(VLOOKUP(F{row},'06算法成本画像'!$A:$K,6,FALSE),0)")
        ws_calc.cell(row, 10, f'=IF(F{row}="",0,D{row}*H{row}*I{row})')
        ws_calc.cell(row, 11, f'=IFERROR(VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,5,FALSE),"")')
        ws_calc.cell(row, 12, f'=IFERROR(VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,9,FALSE),"")')
        ws_calc.cell(row, 13, f"=IFERROR(VLOOKUP(F{row},'06算法成本画像'!$A:$K,10,FALSE),0)")
        ws_calc.cell(row, 14, f'=IF(OR(F{row}="",L{row}="none",L{row}="optional"),0,D{row}*M{row})')
        ws_calc.cell(row, 15, f'=IF(F{row}="",0,IF(COUNTIF($F$2:F{row},F{row})=1,1,0))')
        ws_calc.cell(row, 16, f'=IF(O{row}=1,VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,7,FALSE),0)')
        ws_calc.cell(row, 17, f'=IF(O{row}=1,VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,8,FALSE),0)')
        ws_calc.cell(row, 18, f'=IF(F{row}="","",VLOOKUP(F{row},\'06算法成本画像\'!$A:$K,11,FALSE))')

    # 05 summary
    add_headers(ws_summary, ["项目", "值", "解释"])
    summary_rows = [
        ("安全余量系数", 1.25, "为了避免后期加硬件困难，建议默认至少预留25%。"),
        ("token/API费用口径", "按用量另计，不进入固定成本", "大模型费用作为用量包或按量收费。"),
        ("", "", ""),
        ("总摄像头路数", "=SUM('02摄像头组输入'!C2:C100)", "只按摄像头组输入计路数，不因选择多个算法重复计算摄像头数量。"),
        ("功能选择行数", '=COUNTIF(\'03功能选择输入\'!B2:B200,"<>")', "一组摄像头选多个功能，会产生多行功能选择。"),
        ("唯一功能数量", "=SUM('04计算明细'!O2:O200)", "用于计算研发/适配成本。"),
        ("基础总算力", "=SUM('04计算明细'!J2:J200)", "所有“摄像头组×功能”的算力相加。"),
        ("预留后所需算力", "=B8*B2", "基础总算力×安全余量。"),
        ("预计AI事件量/月", "=SUM('04计算明细'!N2:N200)*30", "只用于估算大模型调用量，不进入固定成本。"),
        ("", "", ""),
        ("推荐硬件ID", "=INDEX('08硬件档位'!A2:A100,MATCH(1,'08硬件档位'!K2:K100,0))", "硬件档位需同时满足预留后算力和总路数。"),
        ("推荐硬件名称", "=VLOOKUP(B12,'08硬件档位'!A:K,2,FALSE)", ""),
        ("集群节点数", '=IF(B12="cluster",CEILING(B9/VLOOKUP("edge-pro",\'08硬件档位\'!A:C,3,FALSE),1),1)', "若推荐为集群方案，则按高性能服务器节点数估算。"),
        ("硬件买断成本", '=IF(B12="cluster",B14*VLOOKUP("edge-pro",\'08硬件档位\'!A:E,5,FALSE),VLOOKUP(B12,\'08硬件档位\'!A:E,5,FALSE))', ""),
        ("硬件月付成本", '=IF(B12="cluster",B14*VLOOKUP("edge-pro",\'08硬件档位\'!A:F,6,FALSE),VLOOKUP(B12,\'08硬件档位\'!A:F,6,FALSE))', ""),
        ("一次性功能研发/适配成本", "=SUM('04计算明细'!P2:P200)", "同一功能只计一次研发/适配成本。"),
        ("月度算法运维成本", "=SUM('04计算明细'!Q2:Q200)", ""),
        ("固定买断合计", "=B15+B17", "硬件买断 + 一次性研发/适配。"),
        ("固定月付合计", "=B16+B18", "硬件月付 + 月度算法运维，不含token/API。"),
        ("token/API费用", "按实际调用量或服务包另计", ""),
    ]
    for row in summary_rows:
        ws_summary.append(row)

    # 09 notes
    add_headers(ws_notes, ["问题", "说明"])
    for row in [
        ("为什么不用一个大表？", "因为客户需求、算法资源画像、硬件容量画像、报价规则是四类不同的数据。拆开后，每个参数都能单独校准。"),
        ("为什么功能研发成本不乘路数？", "通常同一功能的研发/适配是一次性的，路数增加主要影响算力和硬件，不应重复收同一功能研发费。"),
        ("为什么算法算力要乘分辨率？", "分辨率越高，解码、检测和后处理压力越高。4K即使路数不多，也可能吃掉大量硬件余量。"),
        ("为什么token/API不进固定成本？", "大模型不是每秒跑，而是事件候选触发后调用，适合作为用量包或按量计费。"),
        ("最需要校准什么？", "06算法成本画像中的单路算力、研发成本、AI事件率；08硬件档位中的容量和真实成本。"),
    ]:
        ws_notes.append(row)

    # Apply styles.
    for ws in wb.worksheets:
        setup_sheet(ws)

    for ws in [ws_groups, ws_features]:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                style_cell(cell, fill="FFF7DD")

    for row in [15, 16, 17, 18, 19, 20, 21]:
        for col in range(1, 4):
            style_cell(ws_summary.cell(row, col), fill="EAF8F4" if row in [19, 20] else "EAF2FF", bold=row in [19, 20])

    # Validations.
    dv_res = DataValidation(type="list", formula1=f"'07分辨率倍率'!$A$2:$A${1 + len(resolutions)}", allow_blank=True)
    ws_groups.add_data_validation(dv_res)
    dv_res.add("D2:D100")

    dv_group = DataValidation(type="list", formula1="'02摄像头组输入'!$A$2:$A$100", allow_blank=True)
    ws_features.add_data_validation(dv_group)
    dv_group.add("A2:A200")

    dv_feature = DataValidation(type="list", formula1=f"'06算法成本画像'!$A$2:$A${1 + len(feature_costs)}", allow_blank=True)
    ws_features.add_data_validation(dv_feature)
    dv_feature.add("B2:B200")

    # Comments.
    ws_summary["B2"].comment = Comment("这个系数越高，硬件越保守，后期加硬件风险越低。", "Codex")
    ws_feature_cost["F1"].comment = Comment("以1080P单路摄像头为基准的算法算力单位。不是钱，是硬件压力。", "Codex")
    ws_feature_cost["G1"].comment = Comment("功能研发/适配成本，按唯一功能计一次。", "Codex")
    ws_feature_cost["J1"].comment = Comment("用于估算大模型调用量，不直接计入固定成本。", "Codex")

    # Hardware conditional formatting.
    for row in range(2, 2 + len(hardware)):
        ws_hw.conditional_formatting.add(
            f"K{row}",
            FormulaRule(formula=[f"$K{row}=1"], fill=PatternFill("solid", fgColor="EAF8F4")),
        )

    custom_widths = {
        "01逻辑总览": {"A": 24, "B": 96},
        "02摄像头组输入": {"A": 12, "B": 24, "C": 12, "D": 12, "E": 12, "F": 14, "G": 36},
        "03功能选择输入": {"A": 12, "B": 24, "C": 28, "D": 32},
        "04计算明细": {"A": 8, "B": 10, "C": 24, "D": 12, "E": 12, "F": 24, "G": 28, "J": 12, "P": 18, "Q": 18, "R": 42},
        "05报价汇总": {"A": 28, "B": 24, "C": 70},
        "06算法成本画像": {"A": 24, "B": 28, "C": 18, "D": 18, "E": 12, "F": 16, "G": 18, "H": 18, "I": 14, "J": 18, "K": 40},
        "08硬件档位": {"A": 16, "B": 22, "C": 16, "D": 16, "E": 14, "F": 14, "G": 18, "H": 24, "I": 14, "J": 42, "K": 16},
        "09参数说明": {"A": 28, "B": 100},
    }
    for ws_name, widths in custom_widths.items():
        for col, width in widths.items():
            wb[ws_name].column_dimensions[col].width = width

    # Number formats.
    for ws in [ws_summary, ws_calc, ws_feature_cost, ws_hw, ws_res]:
        for row in ws.iter_rows():
            for cell in row:
                if cell.column in [2, 3, 4, 5, 6, 8, 9, 10, 13, 14, 15, 16, 17]:
                    cell.number_format = "#,##0.00"

    wb.active = wb.sheetnames.index("01逻辑总览")
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_workbook())
